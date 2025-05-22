from docx import Document

# 문서 불러오기
doc = Document("sample.docx")

# 문서 내의 첫 번째 표 선택
table = doc.tables[0]

# 예: 1행 2열(0-indexed)의 셀 값을 수정
table.cell(1, 2).text = "새로운 값"

# 저장
doc.save("updated_sample.docx")